# -*- coding: utf-8 -*-
"""Пять листов выгрузки: ДДС, ОСВ, Баланс, Журнал проводок, За балансом (ARCHITECTURE §9).

Цифры на листах обязаны совпадать с экранами: и лист «Баланс», и Главная считают одной
функцией ledger.balances(); ДДС строится из тех же проводок, что и аналитика."""
from __future__ import annotations

import os
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from backend import db, rates
from backend.ledger.balances import balance_map
from backend.ledger.common import money_dict, natural_sign, to_view
from backend.tz import add_days, month_bounds, shift_month, today_dubai
from backend.ledger.tranches import open_rows as open_tranche_rows

MONEY_FORMAT = '# ##0,00;[RED]-# ##0,00'
SHEETS = ("dds", "osv", "balance", "journal", "offbalance")

INFLOW_LABELS = {
    "loan_repaid_to_me": "Возвраты долгов мне",
    "family_withdrawal": "Взято из семейного",
    "loan_taken": "Взято в долг",
}
OUTFLOW_LABELS = {
    "tax_payment": "Уплата налога",
    "loan_given": "Выдано в долг",
    "loan_repaid_by_me": "Возврат долга",
    "family_repayment": "Возврат в семейный",
    "own_contribution_to_family": "Взнос в семейный",
    "paid_for_family_from_own": "Оплата за семью",
}


def period_bounds(period: str | None, date_from: str | None, date_to: str | None) -> tuple[str, str]:
    if period:
        return month_bounds(period)
    if date_from and date_to:
        return date_from, date_to
    return month_bounds(today_dubai()[:7])


def months_between(date_from: str, date_to: str) -> list[str]:
    out, cursor = [], date_from[:7]
    while cursor <= date_to[:7]:
        out.append(cursor)
        cursor = shift_month(cursor, 1)
    return out


def money_account_ids(conn: sqlite3.Connection) -> dict[int, sqlite3.Row]:
    rows = conn.execute(
        "SELECT * FROM accounts WHERE is_money = 1 AND type = 'asset_personal'").fetchall()
    return {int(row["id"]): row for row in rows}


def _line_for(conn: sqlite3.Connection, row: sqlite3.Row, money_ids: set[int],
              names: dict[int, str]) -> tuple[str, str] | None:
    """(«Поступления»|«Выплаты», статья) для одной проводки по денежному счёту."""
    kind = row["kind"]
    amount = int(row["amount"])
    if kind == "opening_balance":
        return None                        # только в «Остаток на начало»
    if kind == "expense":
        other = conn.execute(
            "SELECT a.id, a.name FROM postings p JOIN accounts a ON a.id = p.account_id"
            " WHERE p.entry_id = ? AND a.type = 'expense' LIMIT 1",
            (int(row["entry_id"]),)).fetchone()
        return ("Выплаты", other["name"] if other else "Прочие выплаты")
    if kind == "income":
        other = conn.execute(
            "SELECT a.id, a.name FROM postings p JOIN accounts a ON a.id = p.account_id"
            " WHERE p.entry_id = ? AND a.type = 'income' LIMIT 1",
            (int(row["entry_id"]),)).fetchone()
        return ("Поступления", other["name"] if other else "Прочие приходы")
    if kind == "transfer":
        others = conn.execute(
            "SELECT a.id, a.type, a.is_money FROM postings p JOIN accounts a ON a.id = p.account_id"
            " WHERE p.entry_id = ? AND p.account_id <> ?",
            (int(row["entry_id"]), int(row["account_id"]))).fetchall()
        for other in others:
            if other["type"] == "liability":
                return ("Выплаты", "Погашение кредитки")
        return ("Поступления" if amount > 0 else "Выплаты", "Переводы между своими")
    if kind == "exchange":
        return ("Поступления" if amount > 0 else "Выплаты", "Переоценка/обмен")
    if kind == "balance_correction":
        return ("Поступления" if amount > 0 else "Выплаты", "Корректировка остатков")
    if amount > 0:
        return ("Поступления", INFLOW_LABELS.get(kind, "Прочие приходы"))
    return ("Выплаты", OUTFLOW_LABELS.get(kind, "Прочие выплаты"))


def dds_data(conn: sqlite3.Connection, date_from: str, date_to: str) -> dict:
    """Строки ДДС: каждая проводка по денежному счёту попадает ровно в одну статью (§9)."""
    money = money_account_ids(conn)
    money_ids = set(money)
    names = {int(i): row["name"] for i, row in money.items()}
    months = months_between(date_from, date_to)

    rows = conn.execute(
        "SELECT p.account_id AS account_id, p.amount AS amount, p.currency AS currency,"
        " p.entry_id AS entry_id, e.kind AS kind, e.op_date AS op_date"
        " FROM postings p JOIN entries e ON e.id = p.entry_id"
        " WHERE e.status = 'posted' AND e.sealed = 1 AND e.op_date >= ? AND e.op_date <= ?",
        (date_from, date_to)).fetchall()

    lines: dict[tuple[str, str, str], dict[str, int]] = {}
    opening_in_period: dict[str, int] = {}
    for row in rows:
        if int(row["account_id"]) not in money_ids:
            continue
        currency = row["currency"]
        amount = int(row["amount"])
        if row["kind"] == "opening_balance":
            opening_in_period[currency] = opening_in_period.get(currency, 0) + amount
            continue
        line = _line_for(conn, row, money_ids, names)
        if line is None:
            continue
        section, label = line
        key = (section, label, currency)
        bucket = lines.setdefault(key, {month: 0 for month in months})
        bucket[row["op_date"][:7]] = bucket.get(row["op_date"][:7], 0) + amount

    before = balance_map(conn, as_of=add_days(date_from, -1))
    end = balance_map(conn, as_of=date_to)
    start_totals: dict[str, int] = {}
    end_totals: dict[str, int] = {}
    for (account_id, currency), amount in before.items():
        if account_id in money_ids:
            start_totals[currency] = start_totals.get(currency, 0) + amount
    for currency, amount in opening_in_period.items():
        start_totals[currency] = start_totals.get(currency, 0) + amount
    for (account_id, currency), amount in end.items():
        if account_id in money_ids:
            end_totals[currency] = end_totals.get(currency, 0) + amount

    net: dict[str, int] = {}
    for (section, label, currency), values in lines.items():
        net[currency] = net.get(currency, 0) + sum(values.values())
    return {"months": months, "lines": lines, "start": start_totals, "end": end_totals,
            "net": net}


def osv_data(conn: sqlite3.Connection, date_from: str, date_to: str) -> list[dict]:
    before = balance_map(conn, as_of=add_days(date_from, -1))
    turns = conn.execute(
        "SELECT p.account_id AS account_id, p.currency AS currency,"
        " SUM(CASE WHEN p.amount > 0 THEN p.amount ELSE 0 END) AS dt,"
        " SUM(CASE WHEN p.amount < 0 THEN -p.amount ELSE 0 END) AS kt"
        " FROM postings p JOIN entries e ON e.id = p.entry_id"
        " WHERE e.status = 'posted' AND e.sealed = 1 AND e.op_date >= ? AND e.op_date <= ?"
        " GROUP BY p.account_id, p.currency", (date_from, date_to)).fetchall()
    accounts = {int(row["id"]): row for row in conn.execute("SELECT * FROM accounts").fetchall()}
    keys = set(before) | {(int(r["account_id"]), r["currency"]) for r in turns}
    turn_map = {(int(r["account_id"]), r["currency"]): (int(r["dt"] or 0), int(r["kt"] or 0))
                for r in turns}
    out = []
    for account_id, currency in sorted(keys, key=lambda k: (accounts[k[0]]["sort"], k[0], k[1])):
        account = accounts[account_id]
        start = before.get((account_id, currency), 0)
        dt, kt = turn_map.get((account_id, currency), (0, 0))
        end = start + dt - kt
        if start == 0 and dt == 0 and kt == 0 and end == 0:
            continue
        out.append({"account_id": account_id, "slug": account["slug"], "name": account["name"],
                    "type": account["type"], "currency": currency,
                    "start_dt": max(start, 0), "start_kt": max(-start, 0),
                    "turn_dt": dt, "turn_kt": kt,
                    "end_dt": max(end, 0), "end_kt": max(-end, 0)})
    return out


def balance_data(conn: sqlite3.Connection, date_to: str) -> dict:
    bmap = balance_map(conn, as_of=date_to)
    snapshot = rates.snapshot_at_date(date_to, conn=conn)
    accounts = {int(row["id"]): row for row in conn.execute("SELECT * FROM accounts").fetchall()}
    sections = {"asset_personal": "Активы личные", "asset_family": "Активы семейные",
                "liability": "Обязательства", "equity": "Капитал"}
    rows = []
    totals = {"assets": 0, "liabilities": 0, "net_worth": 0}
    for (account_id, currency), amount in sorted(bmap.items(),
                                                 key=lambda kv: (accounts[kv[0][0]]["sort"], kv[0][0])):
        account = accounts[account_id]
        if account["type"] not in sections or amount == 0:
            continue
        in_rub = to_view(amount, currency, "RUB", snapshot)
        rows.append({"section": sections[account["type"]], "name": account["name"],
                     "currency": currency, "amount": amount * natural_sign(account["type"]),
                     "in_rub": in_rub * natural_sign(account["type"]), "type": account["type"]})
        if account["type"] in ("asset_personal", "asset_family"):
            totals["assets"] += in_rub
        if account["type"] == "liability":
            totals["liabilities"] += -in_rub
        if account["type"] in ("asset_personal", "asset_family", "liability"):
            totals["net_worth"] += in_rub
    from backend.ledger.balances import family_debt
    debt = family_debt(conn, as_of=date_to, snapshot=snapshot, bmap=bmap)
    return {"rows": rows, "totals": totals, "family_debt": debt["debt"], "snapshot": snapshot}


def journal_data(conn: sqlite3.Connection, date_from: str, date_to: str) -> list[dict]:
    rows = conn.execute(
        "SELECT e.id AS entry_id, e.op_date AS op_date, e.kind AS kind, e.description AS descr,"
        " e.source AS source, e.status AS status, e.rates_snapshot AS snap, p.seq AS seq,"
        " p.amount AS amount, p.currency AS currency, p.project AS project,"
        " a.name AS account_name, c.name AS counterparty"
        " FROM entries e JOIN postings p ON p.entry_id = e.id"
        " JOIN accounts a ON a.id = p.account_id"
        " LEFT JOIN counterparties c ON c.id = p.counterparty_id"
        " WHERE e.op_date >= ? AND e.op_date <= ? ORDER BY e.op_date, e.id, p.seq",
        (date_from, date_to)).fetchall()
    return [dict(row) for row in rows]


def offbalance_data(conn: sqlite3.Connection) -> dict:
    tranche_rows = [dict(row) for row in open_tranche_rows(conn, status="all")]
    goals = conn.execute("SELECT * FROM goals ORDER BY id").fetchall()
    goal_rows = []
    for goal in goals:
        saved = conn.execute(
            "SELECT COALESCE(SUM(amount), 0) AS total FROM goal_allocations WHERE goal_id = ?",
            (int(goal["id"]),)).fetchone()["total"]
        goal_rows.append({"name": goal["name"], "target": int(goal["target"]),
                          "saved": int(saved), "currency": goal["currency"],
                          "due_date": goal["due_date"]})
    return {"tranches": tranche_rows, "goals": goal_rows}


# --------------------------------------------------------------------------------------
# сама книга


def _header(sheet, title: str, columns: list[str]) -> None:
    sheet["A1"] = title
    sheet["A1"].font = Font(bold=True, size=13)
    for index, name in enumerate(columns, start=1):
        cell = sheet.cell(row=3, column=index, value=name)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:{get_column_letter(len(columns))}3"
    widths = [34, 10] + [16] * (len(columns) - 2)
    for index, width in enumerate(widths[:len(columns)], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _money(sheet, row: int, column: int, minor: int) -> None:
    cell = sheet.cell(row=row, column=column, value=int(minor) / 100)
    cell.number_format = MONEY_FORMAT


def build(*, period: str | None = None, date_from: str | None = None, date_to: str | None = None,
          sheets: list[str] | None = None, path: str | None = None) -> tuple[str, str]:
    conn = db.get_conn()
    date_from, date_to = period_bounds(period, date_from, date_to)
    wanted = [s for s in (sheets or SHEETS) if s in SHEETS] or list(SHEETS)
    title_period = period or f"{date_from} — {date_to}"

    book = Workbook()
    book.remove(book.active)

    if "dds" in wanted:
        data = dds_data(conn, date_from, date_to)
        sheet = book.create_sheet("ДДС")
        columns = ["Статья", "Валюта"] + data["months"] + ["Итого"]
        _header(sheet, f"Движение денежных средств, {title_period}", columns)
        row_index = 4
        for currency in sorted(set(list(data["start"]) + list(data["end"]) +
                                   [key[2] for key in data["lines"]])):
            sheet.cell(row=row_index, column=1, value="Остаток на начало")
            sheet.cell(row=row_index, column=2, value=currency)
            _money(sheet, row_index, len(columns), data["start"].get(currency, 0))
            row_index += 1
            for section in ("Поступления", "Выплаты"):
                items = sorted([(key, values) for key, values in data["lines"].items()
                                if key[0] == section and key[2] == currency],
                               key=lambda kv: kv[0][1])
                if not items:
                    continue
                sheet.cell(row=row_index, column=1, value=section).font = Font(bold=True)
                row_index += 1
                for (_, label, _), values in items:
                    sheet.cell(row=row_index, column=1, value=label)
                    sheet.cell(row=row_index, column=2, value=currency)
                    total = 0
                    for offset, month in enumerate(data["months"]):
                        value = values.get(month, 0)
                        total += value
                        _money(sheet, row_index, 3 + offset, value)
                    _money(sheet, row_index, len(columns), total)
                    row_index += 1
            sheet.cell(row=row_index, column=1, value="Чистый поток").font = Font(bold=True)
            sheet.cell(row=row_index, column=2, value=currency)
            _money(sheet, row_index, len(columns), data["net"].get(currency, 0))
            row_index += 1
            sheet.cell(row=row_index, column=1, value="Остаток на конец").font = Font(bold=True)
            sheet.cell(row=row_index, column=2, value=currency)
            _money(sheet, row_index, len(columns), data["end"].get(currency, 0))
            row_index += 2

    if "osv" in wanted:
        sheet = book.create_sheet("ОСВ")
        columns = ["Счёт", "Тип", "Валюта", "Сальдо нач. Дт", "Сальдо нач. Кт", "Оборот Дт",
                   "Оборот Кт", "Сальдо кон. Дт", "Сальдо кон. Кт"]
        _header(sheet, f"Оборотно-сальдовая ведомость, {title_period}", columns)
        rows = osv_data(conn, date_from, date_to)
        row_index = 4
        totals: dict[str, list[int]] = {}
        for item in rows:
            sheet.cell(row=row_index, column=1, value=item["name"])
            sheet.cell(row=row_index, column=2, value=item["type"])
            sheet.cell(row=row_index, column=3, value=item["currency"])
            for offset, key in enumerate(("start_dt", "start_kt", "turn_dt", "turn_kt",
                                          "end_dt", "end_kt")):
                _money(sheet, row_index, 4 + offset, item[key])
            bucket = totals.setdefault(item["currency"], [0] * 6)
            for offset, key in enumerate(("start_dt", "start_kt", "turn_dt", "turn_kt",
                                          "end_dt", "end_kt")):
                bucket[offset] += item[key]
            row_index += 1
        for currency, bucket in sorted(totals.items()):
            sheet.cell(row=row_index, column=1, value=f"Итого {currency}").font = Font(bold=True)
            sheet.cell(row=row_index, column=3, value=currency)
            for offset, value in enumerate(bucket):
                _money(sheet, row_index, 4 + offset, value)
            row_index += 1

    if "balance" in wanted:
        sheet = book.create_sheet("Баланс")
        columns = ["Раздел", "Счёт", "Валюта", "Сумма в валюте", "Сумма в ₽"]
        _header(sheet, f"Баланс на {date_to}", columns)
        data = balance_data(conn, date_to)
        row_index = 4
        for item in data["rows"]:
            sheet.cell(row=row_index, column=1, value=item["section"])
            sheet.cell(row=row_index, column=2, value=item["name"])
            sheet.cell(row=row_index, column=3, value=item["currency"])
            _money(sheet, row_index, 4, item["amount"])
            _money(sheet, row_index, 5, item["in_rub"])
            row_index += 1
        for label, value in (("Итого активы", data["totals"]["assets"]),
                             ("Итого обязательства", data["totals"]["liabilities"]),
                             ("Личный итог", data["totals"]["net_worth"])):
            sheet.cell(row=row_index, column=1, value=label).font = Font(bold=True)
            _money(sheet, row_index, 5, value)
            row_index += 1
        sheet.cell(row=row_index, column=1, value="Долг семье, AED")
        _money(sheet, row_index, 4, data["family_debt"]["amount"])

    if "journal" in wanted:
        sheet = book.create_sheet("Журнал проводок")
        columns = ["ID", "Дата", "Тип", "Описание", "Источник ввода", "Статус", "№", "Счёт",
                   "Дт", "Кт", "Валюта", "Контрагент", "Проект", "Курс"]
        _header(sheet, f"Журнал проводок, {title_period}", columns)
        row_index = 4
        totals: dict[str, list[int]] = {}
        for item in journal_data(conn, date_from, date_to):
            amount = int(item["amount"])
            sheet.cell(row=row_index, column=1, value=int(item["entry_id"]))
            sheet.cell(row=row_index, column=2, value=item["op_date"])
            sheet.cell(row=row_index, column=3, value=item["kind"])
            sheet.cell(row=row_index, column=4, value=item["descr"])
            sheet.cell(row=row_index, column=5, value=item["source"])
            sheet.cell(row=row_index, column=6, value=item["status"])
            sheet.cell(row=row_index, column=7, value=int(item["seq"]))
            sheet.cell(row=row_index, column=8, value=item["account_name"])
            if amount > 0:
                _money(sheet, row_index, 9, amount)
            else:
                _money(sheet, row_index, 10, -amount)
            sheet.cell(row=row_index, column=11, value=item["currency"])
            sheet.cell(row=row_index, column=12, value=item["counterparty"])
            sheet.cell(row=row_index, column=13, value=item["project"])
            if item["status"] == "posted":
                bucket = totals.setdefault(item["currency"], [0, 0])
                bucket[0 if amount > 0 else 1] += abs(amount)
            row_index += 1
        for currency, (dt, kt) in sorted(totals.items()):
            sheet.cell(row=row_index, column=8, value=f"Итого {currency}").font = Font(bold=True)
            _money(sheet, row_index, 9, dt)
            _money(sheet, row_index, 10, kt)
            sheet.cell(row=row_index, column=11, value=currency)
            row_index += 1

    if "offbalance" in wanted:
        sheet = book.create_sheet("За балансом")
        columns = ["Дата", "Операция", "Взято AED", "Курс", "Заморожено ₽", "Остаток AED",
                   "Остаток ₽", "Счёт", "Статус"]
        _header(sheet, "За балансом: заморозка и цели", columns)
        row_index = 4
        accounts = {int(row["id"]): row["name"] for row in
                    conn.execute("SELECT id, name FROM accounts").fetchall()}
        data = offbalance_data(conn)
        total_left = 0
        for item in data["tranches"]:
            sheet.cell(row=row_index, column=1, value=item["op_date"])
            sheet.cell(row=row_index, column=2, value=int(item["entry_id"]))
            _money(sheet, row_index, 3, int(item["aed_taken"]))
            sheet.cell(row=row_index, column=4, value=int(item["rate_e6"]) / 1_000_000)
            _money(sheet, row_index, 5, int(item["rub_frozen"]))
            _money(sheet, row_index, 6, int(item["aed_left"]))
            _money(sheet, row_index, 7, int(item["rub_left"]))
            sheet.cell(row=row_index, column=8,
                       value=accounts.get(item["account_id"]) if item["account_id"] else "")
            sheet.cell(row=row_index, column=9, value=item["status"])
            if item["status"] == "open":
                total_left += int(item["rub_left"])
            row_index += 1
        sheet.cell(row=row_index, column=1, value="Итого заморожено").font = Font(bold=True)
        _money(sheet, row_index, 7, total_left)
        row_index += 2
        sheet.cell(row=row_index, column=1, value="Цели").font = Font(bold=True)
        row_index += 1
        for header_index, name in enumerate(["Название", "Цель", "Отложено", "Прогресс %", "Срок"],
                                            start=1):
            sheet.cell(row=row_index, column=header_index, value=name).font = Font(bold=True)
        row_index += 1
        for goal in data["goals"]:
            sheet.cell(row=row_index, column=1, value=goal["name"])
            _money(sheet, row_index, 2, goal["target"])
            _money(sheet, row_index, 3, goal["saved"])
            sheet.cell(row=row_index, column=4,
                       value=round(goal["saved"] * 100 / goal["target"], 1) if goal["target"] else 0)
            sheet.cell(row=row_index, column=5, value=goal["due_date"])
            row_index += 1

    filename = f"finuchet_{period}.xlsx" if period else f"finuchet_{date_from}_{date_to}.xlsx"
    path = path or os.path.join(db.tmp_dir(), filename)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    book.save(path)
    return path, filename
