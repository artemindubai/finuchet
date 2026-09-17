# -*- coding: utf-8 -*-
"""Файл-выписка (PDF под паролем, XLSX, CSV) → те же строки, что у скрина (ТЗ 5.3).

Дальше — общий путь сверки: ledger.reconcile → карточка §7.4 → apply_reconcile.
"""
from __future__ import annotations

import csv
import io
import logging
import os
import re
from typing import Any

from backend.bot.cards import Reply
from backend.pipeline import screenshot as screenshot_pipeline

log = logging.getLogger("finuchet.pipeline.statement")

_DATE_PATTERNS = (
    (re.compile(r"^(\d{4})-(\d{2})-(\d{2})"), (1, 2, 3)),
    (re.compile(r"^(\d{2})[./](\d{2})[./](\d{4})"), (3, 2, 1)),
)


def pdf_password() -> str | None:
    return (os.getenv("MASHREQ_PDF_PASSWORD") or "").strip() or None


def pdf_to_text(blob: bytes, *, password: str | None = None) -> str:
    """Текст PDF. Пароль снимается до отправки в модель; расшифрованная копия на диск не пишется."""
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover
        log.warning("нет pypdf — PDF не разобрать")
        return ""
    try:
        reader = PdfReader(io.BytesIO(blob))
        if reader.is_encrypted:
            if not password:
                log.warning("PDF под паролем, а пароль не задан (MASHREQ_PDF_PASSWORD)")
                return ""
            reader.decrypt(password)
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as exc:
        log.warning("PDF не читается: %s", exc)
        return ""


def pdf_to_images(blob: bytes, *, password: str | None = None, limit: int = 8) -> list[tuple[bytes, str]]:
    """Скан без текстового слоя — страницы картинками (нужен pypdfium2)."""
    try:
        import pypdfium2 as pdfium
    except ImportError:
        log.warning("нет pypdfium2 — скан PDF не превратить в картинки")
        return []
    try:
        doc = pdfium.PdfDocument(io.BytesIO(blob), password=password)
        out: list[tuple[bytes, str]] = []
        for index in range(min(len(doc), limit)):
            bitmap = doc[index].render(scale=2)
            buf = io.BytesIO()
            bitmap.to_pil().save(buf, format="PNG")
            out.append((buf.getvalue(), "image/png"))
        return out
    except Exception as exc:
        log.warning("страницы PDF не отрисовались: %s", exc)
        return []


def rows_from_table(blob: bytes, filename: str) -> list[dict[str, Any]]:
    """XLSX/CSV → строки формата §6.3 (дата, сумма со знаком, текст). Модель не зовём."""
    name = (filename or "").lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        return _rows_from_csv(blob)
    return _rows_from_xlsx(blob)


def _rows_from_csv(blob: bytes) -> list[dict[str, Any]]:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = blob.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        return []
    sample = text[:4096]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    return _rows_from_cells(list(csv.reader(io.StringIO(text), delimiter=delimiter)))


def _rows_from_xlsx(blob: bytes) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover
        return []
    try:
        book = load_workbook(io.BytesIO(blob), data_only=True, read_only=True)
    except Exception as exc:
        log.warning("XLSX не читается: %s", exc)
        return []
    sheet = book[book.sheetnames[0]]
    return _rows_from_cells([[c for c in row] for row in sheet.iter_rows(values_only=True)])


def _rows_from_cells(table: list[list[Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for raw in table:
        cells = ["" if c is None else str(c).strip() for c in raw]
        if not any(cells):
            continue
        op_date = _find_date(cells)
        amount = _find_amount(cells)
        if op_date is None or amount is None:
            continue
        text = " ".join(c for c in cells if c and c != str(amount))[:200]
        rows.append({"op_date": op_date, "amount": amount, "currency": None,
                     "raw_text": text, "suggested_category": None, "suggested_kind": None})
    return rows


def _find_date(cells: list[str]) -> str | None:
    for cell in cells:
        for pattern, order in _DATE_PATTERNS:
            match = pattern.match(cell)
            if match:
                y, m, d = (match.group(i) for i in order)
                return f"{y}-{m}-{d}"
    return None


def _find_amount(cells: list[str]) -> int | None:
    for cell in reversed(cells):
        value = cell.replace(" ", "").replace(" ", "").replace(",", ".")
        if not re.fullmatch(r"[-+]?\d+(\.\d{1,2})?", value or ""):
            continue
        try:
            return int(round(float(value) * 100))
        except ValueError:
            continue
    return None


async def handle_statement(
    chat_id: int,
    blob: bytes,
    filename: str,
    *,
    account_hint: str | None = None,
) -> list[Reply]:
    """Выписка → сверка. PDF идёт через зрение, таблицы разбираются на месте."""
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return await screenshot_pipeline.handle_parsed(
            chat_id,
            await _parse_pdf(blob),
            source="statement",
            account_hint=account_hint,
        )
    rows = rows_from_table(blob, filename)
    if not rows:
        return [Reply("Не разобрал файл. Пришлите выписку PDF, XLSX или CSV.")]
    parsed = {"account_guess": account_hint, "currency": None, "balance": None,
              "statement_period": None, "rows": rows}
    return await screenshot_pipeline.handle_parsed(chat_id, parsed, source="statement",
                                                   account_hint=account_hint)


async def _parse_pdf(blob: bytes) -> dict[str, Any]:
    from backend import llm

    return await llm.parse_statement(blob, password=pdf_password(), reference=llm.reference_text())
