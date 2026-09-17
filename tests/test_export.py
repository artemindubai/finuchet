# -*- coding: utf-8 -*-
"""Выгрузка XLSX: ДДС сходится с остатками, ОСВ — с двойной записью (ARCHITECTURE §9, ТЗ 11.5)."""
import os

from openpyxl import load_workbook

from backend import db, ledger
from backend.export import xlsx
from backend.tz import month_bounds, today_dubai
from tests.conftest import base_state, post

MONTH = today_dubai()[:7]
DATE_FROM, DATE_TO = month_bounds(MONTH)


def busy_month():
    base_state()
    post(kind="income", amount=9_000_000, currency="RUB", to_account="tbank",
         income_source="inc_projects", project="Альфа")
    post(kind="expense", amount=235_000, currency="RUB", from_account="tbank",
         category="exp_groceries")
    post(kind="expense", amount=4_300, currency="AED", from_account="mashreq_personal",
         category="exp_cafe")
    post(kind="transfer", amount=2_000_000, currency="RUB", from_account="tbank",
         to_account="sber")
    post(kind="exchange", amount=5_000_000, currency="RUB", from_account="tbank",
         amount_to=51_230, currency_to="USDT", to_account="bybit")
    post(kind="family_withdrawal", amount=50_000, currency="AED", from_account="mashreq_family",
         to_account="mashreq_personal", rate_e6=24_500_000, rate_source="voice")
    post(kind="family_repayment", amount=20_000, currency="AED", from_account="mashreq_personal",
         to_account="mashreq_family")
    post(kind="loan_given", amount=30_000, currency="AED", from_account="mashreq_personal",
         counterparty="Маша")
    post(kind="expense", amount=240_000, currency="RUB", from_account="sber_credit",
         category="exp_clothes")
    post(kind="tax_payment", amount=1_836_000, currency="RUB", from_account="tbank")


def test_dds_net_flow_equals_balance_delta():
    busy_month()
    conn = db.get_conn()
    data = xlsx.dds_data(conn, DATE_FROM, DATE_TO)
    for currency in set(list(data["start"]) + list(data["end"]) + list(data["net"])):
        start = data["start"].get(currency, 0)
        net = data["net"].get(currency, 0)
        end = data["end"].get(currency, 0)
        assert start + net == end, f"ДДС не сходится по {currency}"


def test_tax_payment_present_in_dds():
    busy_month()
    conn = db.get_conn()
    data = xlsx.dds_data(conn, DATE_FROM, DATE_TO)
    tax_line = next(((key, values) for key, values in data["lines"].items()
                     if key[1] == "Уплата налога"), None)
    assert tax_line is not None
    assert sum(tax_line[1].values()) == -1_836_000
    assert tax_line[0][0] == "Выплаты"


def test_osv_double_entry_balances():
    busy_month()
    conn = db.get_conn()
    rows = xlsx.osv_data(conn, DATE_FROM, DATE_TO)
    by_currency: dict[str, list[int]] = {}
    for row in rows:
        bucket = by_currency.setdefault(row["currency"], [0, 0, 0, 0, 0, 0])
        for index, key in enumerate(("start_dt", "start_kt", "turn_dt", "turn_kt",
                                     "end_dt", "end_kt")):
            bucket[index] += row[key]
    for currency, bucket in by_currency.items():
        assert bucket[0] == bucket[1], f"начальное сальдо не сходится по {currency}"
        assert bucket[2] == bucket[3], f"обороты не сходятся по {currency}"
        assert bucket[4] == bucket[5], f"конечное сальдо не сходится по {currency}"


def test_xlsx_totals_match_api():
    busy_month()
    conn = db.get_conn()
    balance_sheet = xlsx.balance_data(conn, DATE_TO)
    home = ledger.balances()
    assert balance_sheet["totals"]["net_worth"] == home["net_worth"]["amount"]
    assert balance_sheet["family_debt"] == home["family"]["debt"]

    data = xlsx.dds_data(conn, DATE_FROM, DATE_TO)
    analytics = ledger.analytics_month(MONTH)
    groceries = next(values for key, values in data["lines"].items() if key[1] == "Продукты")
    from_analytics = next(item for item in analytics["categories"] if item["name"] == "Продукты")
    assert -sum(groceries.values()) == from_analytics["amount"]["amount"]

    path, filename = xlsx.build(period=MONTH)
    assert os.path.exists(path) and filename == f"finuchet_{MONTH}.xlsx"
    book = load_workbook(path)
    assert book.sheetnames == ["ДДС", "ОСВ", "Баланс", "Журнал проводок", "За балансом"]
    dds_sheet = book["ДДС"]
    labels = [dds_sheet.cell(row=row, column=1).value for row in range(4, dds_sheet.max_row + 1)]
    assert "Остаток на начало" in labels and "Остаток на конец" in labels
    os.remove(path)


def test_export_sheets_subset():
    busy_month()
    path, _ = xlsx.build(period=MONTH, sheets=["journal"])
    book = load_workbook(path)
    assert book.sheetnames == ["Журнал проводок"]
    os.remove(path)
