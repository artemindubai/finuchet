# -*- coding: utf-8 -*-
"""Адверсариальные проверки сверки, выгрузки, параллельной записи, копий и архива счетов."""
import glob
import os
import threading

import pytest

from backend import db, ledger
from backend.contracts import LedgerError, ReconcileDecision, ScreenBatch, ScreenRow
from tests.conftest import balance, base_state, post

RATE = 24_500_000


# ---------------------------------------------------------------------------------------
# 8. сверка


def test_two_voice_entries_one_screen_row_gives_question_not_duplicate():
    """Два кофе по 15 AED занесены голосом, на скрине одна строка → вопрос, а не третий кофе."""
    base_state()
    post(kind="expense", amount=1_500, currency="AED", from_account="mashreq_personal",
         category="exp_cafe", op_date="2026-09-16")
    post(kind="expense", amount=1_500, currency="AED", from_account="mashreq_personal",
         category="exp_cafe", op_date="2026-09-16")

    batch = ScreenBatch(account="mashreq_personal", currency="AED", rows=[
        ScreenRow(amount=-1_500, currency="AED", op_date="2026-09-16", raw_text="COFFEE SHOP")])
    result = ledger.reconcile(batch)
    assert len(result["new"]) == 0, "лишней траты быть не должно"
    assert len(result["ambiguous"]) == 1, "два свободных точных кандидата → вопрос (§4.5 п. 3)"
    assert len(result["ambiguous"][0]["candidates"]) == 2


def test_zero_amount_row_does_not_half_apply_the_batch():
    """Строка с суммой 0 не должна ронять пачку, уже применив часть решений.

    Либо весь вызов применяется, либо ничего: иначе интерфейс покажет ошибку,
    а часть операций уже в журнале."""
    base_state()
    batch = ScreenBatch(account="mashreq_personal", currency="AED", rows=[
        ScreenRow(amount=-4_300, currency="AED", op_date="2026-09-17", raw_text="CAFE"),
        ScreenRow(amount=0, currency="AED", op_date="2026-09-17", raw_text="FEE 0.00")])
    result = ledger.reconcile(batch)
    assert len(result["new"]) == 2
    before = int(db.get_conn().execute(
        "SELECT COUNT(*) AS n FROM entries WHERE status = 'posted'").fetchone()["n"])

    decisions = [ReconcileDecision(inbox_id=item["id"], decision="new")
                 for item in result["new"]]
    with pytest.raises(LedgerError) as exc:
        ledger.apply_reconcile(decisions, actor="test")
    assert exc.value.code == "bad_request"
    after = int(db.get_conn().execute(
        "SELECT COUNT(*) AS n FROM entries WHERE status = 'posted'").fetchone()["n"])
    assert after == before, "часть пачки уже записана, хотя вызов завершился ошибкой"


def test_align_recomputes_diff_at_apply_time():
    base_state()
    batch = ScreenBatch(account="tbank", currency="RUB", balance=13_200_000 - 500_000, rows=[
        ScreenRow(amount=-500_000, currency="RUB", op_date="2026-09-17", raw_text="MARKET")])
    result = ledger.reconcile(batch)
    assert result["balance_check"]["diff"]["amount"] == -500_000
    ledger.apply_reconcile([ReconcileDecision(inbox_id=result["new"][0]["id"], decision="new")],
                           actor="test")
    from backend.ledger.reconcile import align
    aligned = align(result["batch_id"], actor="test")
    assert aligned["entry_id"] is None
    assert [w["code"] for w in aligned["warnings"]] == ["already_aligned"]
    assert balance("tbank", "RUB") == 13_200_000 - 500_000
    assert ledger.check_invariants() == []


# ---------------------------------------------------------------------------------------
# 9. выгрузка XLSX


def build_book(month: str):
    from backend.export import xlsx
    path, filename = xlsx.build(period=month)
    from openpyxl import load_workbook
    return load_workbook(path), path


def test_xlsx_all_currencies_dds_osv_balance_match():
    """Месяц с операциями во всех валютах: ДДС = Δ остатков, ОСВ ΣДт=ΣКт, Баланс = Главная."""
    base_state()
    from backend.tz import month_bounds, today_dubai
    month = today_dubai()[:7]
    date_from, date_to = month_bounds(month)

    post(kind="expense", amount=4_300, currency="AED", from_account="mashreq_personal",
         category="exp_cafe")
    post(kind="income", amount=9_000_000, currency="RUB", to_account="tbank",
         income_source="inc_projects", project="Альфа")
    post(kind="exchange", amount=987_654, currency="RUB", from_account="tbank",
         amount_to=101_79, currency_to="USDT", to_account="bybit")
    post(kind="family_withdrawal", amount=50_000, currency="AED", from_account="mashreq_family",
         to_account="mashreq_personal", rate_e6=RATE, rate_source="voice")
    post(kind="family_repayment", amount=20_000, currency="AED", from_account="mashreq_personal",
         to_account="mashreq_family")
    post(kind="tax_payment", amount=300_000, currency="RUB", from_account="tbank")
    post(kind="loan_given", amount=30_000, currency="AED", from_account="mashreq_personal",
         counterparty="Маша")

    from backend.export import xlsx
    conn = db.get_conn()
    from backend.tz import add_days
    from backend.ledger.balances import balance_map

    data = xlsx.dds_data(conn, date_from, date_to)
    money_ids = set(xlsx.money_account_ids(conn))
    before = balance_map(conn, as_of=add_days(date_from, -1))
    after = balance_map(conn, as_of=date_to)
    for currency in set(list(data["start"]) + list(data["end"])):
        delta = (sum(v for (a, c), v in after.items() if a in money_ids and c == currency)
                 - sum(v for (a, c), v in before.items() if a in money_ids and c == currency))
        assert data["start"].get(currency, 0) + data["net"].get(currency, 0) \
            == data["end"].get(currency, 0), f"ДДС не сходится по {currency}"
        # чистый поток без начальных остатков периода
        opening = data["start"].get(currency, 0) - sum(
            v for (a, c), v in before.items() if a in money_ids and c == currency)
        assert data["net"].get(currency, 0) + opening == delta, f"Δ остатков ≠ потоку по {currency}"

    osv = xlsx.osv_data(conn, date_from, date_to)
    by_currency: dict[str, list[int]] = {}
    for item in osv:
        bucket = by_currency.setdefault(item["currency"], [0, 0, 0, 0])
        bucket[0] += item["turn_dt"]
        bucket[1] += item["turn_kt"]
        bucket[2] += item["end_dt"]
        bucket[3] += item["end_kt"]
    for currency, (dt, kt, end_dt, end_kt) in by_currency.items():
        assert dt == kt, f"ОСВ: обороты не сошлись по {currency}: Дт {dt} ≠ Кт {kt}"
        assert end_dt == end_kt, f"ОСВ: конечное сальдо не сошлось по {currency}"

    sheet_balance = xlsx.balance_data(conn, date_to)
    home = ledger.balances()
    assert abs(sheet_balance["totals"]["net_worth"] - home["net_worth"]["amount"]) <= 2, (
        f"лист «Баланс» {sheet_balance['totals']['net_worth']} против Главной "
        f"{home['net_worth']['amount']}")
    assert sheet_balance["family_debt"]["amount"] == home["family"]["debt"]["amount"]

    book, path = build_book(month)
    assert set(book.sheetnames) == {"ДДС", "ОСВ", "Баланс", "Журнал проводок", "За балансом"}
    off = book["За балансом"]
    frozen_rows = [r for r in off.iter_rows(min_row=4, values_only=True) if r[0] == "Итого заморожено"]
    assert frozen_rows, "лист «За балансом» без итога заморозки"
    assert round(frozen_rows[0][6] * 100) == home["frozen_rub"]["amount"]
    assert os.path.getsize(path) > 0


# ---------------------------------------------------------------------------------------
# 12. параллельная запись


def test_parallel_post_operations_keep_invariants():
    """Два потока пишут одновременно: ни одной незапечатанной операции, инварианты целы."""
    base_state()
    errors: list[Exception] = []
    start = threading.Barrier(2)

    def worker(tag: str):
        try:
            start.wait(timeout=10)
            for _ in range(10):
                post(kind="expense", amount=100, currency="AED",
                     from_account="mashreq_personal", category="exp_cafe", description=tag)
        except Exception as exc:                        # noqa: BLE001
            errors.append(exc)
        finally:
            db.close_conn()

    threads = [threading.Thread(target=worker, args=(name,)) for name in ("a", "b")]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join(timeout=60)

    assert not errors, errors
    conn = db.get_conn()
    written = int(conn.execute(
        "SELECT COUNT(*) AS n FROM entries WHERE kind = 'expense'").fetchone()["n"])
    assert written == 20
    unsealed = int(conn.execute(
        "SELECT COUNT(*) AS n FROM entries WHERE sealed = 0").fetchone()["n"])
    assert unsealed == 0
    assert balance("mashreq_personal", "AED") == 184_250 - 20 * 100
    assert ledger.check_invariants() == []


# ---------------------------------------------------------------------------------------
# 13. копии базы


def test_backup_creates_file_and_keeps_only_seven():
    from backend import backup
    base_state()
    path = backup.backup_once()
    assert path and os.path.exists(path)
    row = db.get_conn().execute("SELECT * FROM backups ORDER BY id DESC LIMIT 1").fetchone()
    assert int(row["ok"]) == 1 and int(row["bytes"]) > 0

    directory = db.backup_dir()
    for day in range(1, 9):
        open(os.path.join(directory, f"finuchet-2026-01-{day:02d}.db"), "wb").close()
    backup._rotate()
    left = sorted(os.path.basename(p) for p in glob.glob(os.path.join(directory, "finuchet-*.db")))
    assert len(left) == backup.KEEP, left
    assert "finuchet-2026-01-01.db" not in left, "самая старая копия обязана вытесняться"


# ---------------------------------------------------------------------------------------
# 14. архивные счета


def test_family_account_in_other_currency_fails_cleanly(client):
    """Семейный счёт не в AED: Семейный фонд ведётся в AED, ядро обязано сказать это словами."""
    base_state()
    created = client.post("/api/accounts", json={
        "name": "Семейный рублёвый", "type": "asset_family", "currency": "RUB",
        "is_money": True, "group_key": "FAMILY"})
    assert created.status_code == 200, created.text
    response = client.post("/api/opening-balances", json={
        "items": [{"account": "semeynyy_rublevyy", "amount": 100_000, "currency": "RUB"}]})
    assert response.status_code == 400, response.text
    assert response.json()["error"]["code"] == "currency_mismatch"


def test_archive_account_with_balance_warns(client):
    base_state()
    conn = db.get_conn()
    account_id = int(conn.execute("SELECT id FROM accounts WHERE slug = 'sber'").fetchone()["id"])
    response = client.post(f"/api/accounts/{account_id}/archive")
    assert response.status_code == 200, response.text
    codes = [w["code"] for w in response.json().get("warnings", [])]
    assert "archived_with_balance" in codes, response.json().get("warnings")


def test_posting_to_archived_account_is_rejected(client):
    """Архив = счётом больше не пользуются; новая операция на него — ошибка, а не тихая запись."""
    base_state()
    conn = db.get_conn()
    account_id = int(conn.execute("SELECT id FROM accounts WHERE slug = 'sber'").fetchone()["id"])
    client.post(f"/api/accounts/{account_id}/archive")
    response = client.post("/api/entries", json={
        "kind": "expense", "amount": 100_000, "currency": "RUB",
        "from_account": "sber", "category": "exp_cafe"})
    assert response.status_code == 400, response.text
